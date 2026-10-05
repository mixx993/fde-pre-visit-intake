"""Regenerate fde-pre-visit-intake/template/intake_form.xlsx and fields.json.

Requires openpyxl. Usage: python3 tools/build_template.py [output.xlsx]
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

F = "微软雅黑"
def font(size=10.5, bold=False, color="222222", italic=False):
    return Font(name=F, size=size, bold=bold, color=color, italic=italic)
FILL_IN = PatternFill("solid", fgColor="FFF6CC")     # 待填写
FILL_HEAD = PatternFill("solid", fgColor="1F3A5F")
FILL_SEC = PatternFill("solid", fgColor="E8EEF5")
FILL_EX = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="C8CED6")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()
FIELDS = []
def reg(ws, row, col, fid, label=""):
    from openpyxl.utils import get_column_letter
    FIELDS.append({"id": fid, "sheet": ws.title, "sheet_index": wb.sheetnames.index(ws.title) + 1,
                   "cell": f"{get_column_letter(col)}{row}", "label": label})

def title(ws, text, sub, span):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    ws.cell(1, 1, text).font = font(16, True, "1F3A5F")
    ws.row_dimensions[1].height = 30
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=span)
    c = ws.cell(2, 1, sub); c.font = font(10, color="666666"); c.alignment = WRAP
    ws.row_dimensions[2].height = 36

def section(ws, r, text, span):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=span)
    c = ws.cell(r, 1, text); c.font = font(11.5, True, "1F3A5F"); c.fill = FILL_SEC
    for col in range(1, span + 1):
        ws.cell(r, col).fill = FILL_SEC
    ws.row_dimensions[r].height = 22

def header(ws, r, names):
    for i, n in enumerate(names, 1):
        c = ws.cell(r, i, n); c.font = font(10.5, True, "FFFFFF"); c.fill = FILL_HEAD
        c.alignment = CENTER; c.border = BOX
    ws.row_dimensions[r].height = 30

def qa(ws, r, q, ex, h=34, fid=None):
    if fid: reg(ws, r, 2, fid, q)
    a = ws.cell(r, 1, q); a.font = font(); a.alignment = WRAP; a.border = BOX
    b = ws.cell(r, 2); b.fill = FILL_IN; b.alignment = WRAP; b.border = BOX; b.font = font()
    c = ws.cell(r, 3, ex); c.font = font(10, color="888888", italic=True); c.alignment = WRAP; c.border = BOX; c.fill = FILL_EX
    ws.row_dimensions[r].height = h

def example_row(ws, r, vals):
    for i, v in enumerate(vals, 1):
        c = ws.cell(r, i, v); c.font = font(10, color="888888", italic=True)
        c.fill = FILL_EX; c.alignment = WRAP; c.border = BOX
    ws.row_dimensions[r].height = 34

def input_rows(ws, r0, n, ncol, h=30, fixed=None):
    for r in range(r0, r0 + n):
        for col in range(1, ncol + 1):
            c = ws.cell(r, col); c.fill = FILL_IN; c.border = BOX; c.alignment = WRAP; c.font = font()
        if fixed:
            for col, v in fixed(r - r0).items():
                ws.cell(r, col, v).font = font()
        ws.row_dimensions[r].height = h

def setup(ws, landscape=True):
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

# ---------- 1 填写说明 ----------
ws = wb.active; ws.title = "填写说明"; setup(ws, False)
ws.column_dimensions["A"].width = 4; ws.column_dimensions["B"].width = 92
title(ws, "上门前信息表", "感谢配合！这份表帮助我们在上门前了解您的业务和电脑环境，让上门时间都用在真正需要现场做的事上。", 2)
rows = [
    ("怎么填", None),
    ("1", "共 5 张表：基本情况、需求与流程、电脑与软件、文件清单、特殊设备。第 5 张只有用到工控机、触摸屏、收银机、平板、PLC 这类设备时才填。大约 15–20 分钟。"),
    ("2", "只填黄色格子。灰色斜体是示例，照着格式填就行，不用删。"),
    ("3", "能截图、录屏的就不用打字：在表里写上截图或录屏的文件名即可。"),
    ("4", "不确定、不知道的地方空着，我们上门时再一起确认。"),
    ("文件怎么发", None),
    ("1", "截图、录屏、样例文件请用网盘或 U 盘给我们，不要用微信发：微信会压缩视频和图片，规格会变。"),
    ("2", "录屏请从头到尾录一遍员工平时真实的操作，不用剪辑，有点慢也没关系。"),
    ("请注意", None),
    ("1", "我们不需要、也请不要提供任何账号的密码。"),
    ("2", "账号名称、商品、客户信息可以用代号代替，比如“1 号账号”。"),
    ("3", "填好后连同文件一起发给对接人。有问题随时联系。"),
]
r = 4
for k, v in rows:
    if v is None:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        c = ws.cell(r, 1, k); c.font = font(11.5, True, "1F3A5F"); c.fill = FILL_SEC; ws.cell(r, 2).fill = FILL_SEC
        ws.row_dimensions[r].height = 22
    else:
        ws.cell(r, 1, k).font = font(color="888888"); ws.cell(r, 1).alignment = Alignment(horizontal="center", vertical="top")
        c = ws.cell(r, 2, v); c.font = font(); c.alignment = WRAP
        ws.row_dimensions[r].height = 22 if len(v) < 46 else 36
    r += 1
r += 1
ws.cell(r, 2, "图例").font = font(10.5, True); r += 1
c = ws.cell(r, 2, "黄色格子：请您填写"); c.fill = FILL_IN; c.font = font(); c.border = BOX; r += 1
c = ws.cell(r, 2, "灰色斜体：填写示例，仅供参考"); c.fill = FILL_EX; c.font = font(10, color="888888", italic=True); c.border = BOX

# ---------- 2 基本情况 ----------
ws = wb.create_sheet("1 基本情况"); setup(ws, False)
for col, w in zip("ABC", (34, 46, 40)): ws.column_dimensions[col].width = w
title(ws, "一、基本情况", "团队、业务量、您的期望和限制。只填黄色格子，右边灰色是示例。", 3)
header(ws, 3, ["问题", "请填写", "示例"])
r = 4
BASIC_IDS = ['team_groups', 'decision_maker', 'contact', 'staff_involved', 'accounts', 'daily_output', 'daily_publish', 'peak_time', 'hope_save', 'most_painful', 'worry', 'cooperation_mode', 'cannot_touch', 'no_disturb', 'test_accounts']
bi = iter(BASIC_IDS)
blocks = [
    ("团队和对接", [
        ("公司分几个组？每组几个人？", "运营 5 人、剪辑 2 人，没有分组"),
        ("谁做最终决定？", "王总"),
        ("日常和我们对接的是谁？怎么联系？", "运营组长小李，微信"),
        ("哪些员工在做这次要改进的工作？", "运营 5 人负责发布，剪辑 2 人负责出片"),
    ]),
    ("业务量", [
        ("一共有多少个账号？", "约 20 个"),
        ("每天大约产出多少条内容？", "每天剪约 40 条"),
        ("每天大约发布多少条？", "每个号每天 3 条，共约 60 条"),
        ("一天里最忙的时段？", "晚上 7–9 点集中发布"),
    ]),
    ("期望", [
        ("做成以后，您最希望省掉哪件事？", "每天手动一条条上传和挂商品"),
        ("现在最花时间、最让人头疼的是哪一步？", "上传要等很久，号多了容易漏发"),
        ("对这次合作，您最担心什么？", "账号被平台限流；员工学不会新工具"),
        ("希望我们做一次性工具，还是长期帮忙维护？", "先做一个能用的，后面再看"),
    ]),
    ("限制", [
        ("哪些电脑或账号不能动？", "财务电脑不能动；主力号不要拿来测试"),
        ("什么时间不方便打扰？", "每天晚上 7–9 点发布高峰"),
        ("可以拿来测试的账号有哪些？（写代号即可）", "测试号 A、测试号 B"),
    ]),
]
for name, qs in blocks:
    section(ws, r, name, 3); r += 1
    for q, ex in qs:
        qa(ws, r, q, ex, fid="basic." + next(bi)); r += 1
ws.freeze_panes = "A4"


# ---------- 3 需求与流程 ----------
ws = wb.create_sheet("2 需求与流程"); setup(ws)
cols = ["序号", "这一步做什么", "谁来做", "用什么软件", "每次大约多久", "每天做几次", "这一步最麻烦的地方", "录屏 / 截图文件名"]
widths = [7, 36, 12, 16, 13, 11, 30, 24]
for i, w in enumerate(widths): ws.column_dimensions[chr(65 + i)].width = w
title(ws, "二、需求与现在的做法（SOP）", "每个需求填一块：先写这个需求想解决什么，再按顺序写现在一步步怎么做（谁、用什么、多久）。有现成的 SOP 文档也可以直接发给我们，在这里写文件名。", 8)
r = 4
for n in range(1, 4):
    section(ws, r, f"需求 {n}", 8); r += 1
    for key, q, ex in [("name", "需求名称", "自动发布视频"),
                  ("problem", "想解决什么问题？", "员工每天手动发几百条，太费人"),
                  ("sop_doc", "现成的 SOP 文档（写文件名，没有就空着）", "发布流程.docx")]:
        reg(ws, r, 3, f"needs.{n-1}.{key}", q)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
        ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=8)
        a = ws.cell(r, 1, q); a.font = font(10.5, True); a.alignment = Alignment(vertical="center", wrap_text=True)
        for col in range(1, 9): ws.cell(r, col).border = BOX
        for col in range(3, 7): ws.cell(r, col).fill = FILL_IN
        ws.cell(r, 3).alignment = WRAP; ws.cell(r, 3).font = font()
        for col in range(7, 9): ws.cell(r, col).fill = FILL_EX
        if n == 1:
            c = ws.cell(r, 7, "示例：" + ex); c.font = font(10, color="888888", italic=True); c.alignment = WRAP
        ws.row_dimensions[r].height = 30
        r += 1
    header(ws, r, cols); r += 1
    if n == 1:
        example_row(ws, r, ["示例", "打开发布页，上传视频，填描述，挂商品，保存草稿", "运营", "多账号发布工具", "约 1 分钟", "每号 3 次", "上传要等，等的时候干不了别的", "发布录屏.mp4"]); r += 1
    input_rows(ws, r, 8, 8, fixed=lambda i: {1: i + 1})
    for rr in range(r, r + 8):
        ws.cell(rr, 1).alignment = CENTER
        for ci, key in enumerate(["what", "who", "software", "duration", "per_day", "pain", "file"], 2):
            reg(ws, rr, ci, f"needs.{n-1}.steps.{rr-r}.{key}", cols[ci-1])
    r += 9
ws.freeze_panes = "A4"

# ---------- 4 电脑与软件 ----------
ws = wb.create_sheet("3 电脑与软件"); setup(ws)
widths = [20, 18, 18, 18, 18, 30]
for i, w in enumerate(widths): ws.column_dimensions[chr(65 + i)].width = w
title(ws, "三、电脑、软件和网络", "电脑配置可以直接截“设置 → 系统 → 系统信息”页；软件版本截软件里的“关于”页。写上截图文件名就行，其他格子能填多少填多少。", 6)
r = 4
section(ws, r, "电脑（做这些工作的每台电脑一行）", 6); r += 1
header(ws, r, ["这台电脑的叫法 / 用途", "Windows 版本", "处理器", "内存", "显卡", "“系统信息”截图文件名"]); r += 1
example_row(ws, r, ["发布机 1（运营组用）", "Windows 10 专业版", "i5-12400", "16 GB", "集成显卡", "发布机1-系统信息.png"]); r += 1
input_rows(ws, r, 6, 6)
for rr in range(r, r + 6):
    for ci, key in enumerate(["name", "os", "cpu", "ram", "gpu", "screenshot"], 1):
        reg(ws, rr, ci, f"computers.{rr-r}.{key}")
r += 7

section(ws, r, "软件（和这次需求有关的软件，每个一行）", 6); r += 1
header(ws, r, ["软件名称", "版本号", "怎么装的", "是否付费 / 会员", "用在哪个需求", "“关于”页截图文件名"]); r += 1
example_row(ws, r, ["剪映专业版", "5.9", "官网下载安装", "是", "需求 2 剪辑", "剪映-关于.png"]); r += 1
sw0 = r
input_rows(ws, r, 8, 6)
for rr in range(r, r + 8):
    for ci, key in enumerate(["name", "version", "install", "paid", "need", "screenshot"], 1):
        reg(ws, rr, ci, f"software.{rr-r}.{key}")
r += 9
dv1 = DataValidation(type="list", formula1='"官网下载安装,别人发的安装包,解压直接运行,不清楚"', allow_blank=True)
dv2 = DataValidation(type="list", formula1='"是,否,不清楚"', allow_blank=True)
ws.add_data_validation(dv1); ws.add_data_validation(dv2)
dv1.add(f"C{sw0}:C{sw0+7}"); dv2.add(f"D{sw0}:D{sw0+7}")

section(ws, r, "网络和存储", 6); r += 1
header(ws, r, ["问题", "", "请填写", "", "示例", ""]); 
for a, b in ((1, 2), (3, 4), (5, 6)): ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
r += 1
for key, q, ex in [("upload", "宽带上行速度（可截测速图）", "上行 50 Mbps，测速.png"),
              ("nas", "有没有共享盘或 NAS？", "没有，各存各的电脑"),
              ("storage", "成片和素材平时存在哪？", "每台电脑 D 盘“成片”文件夹，按日期分"),
              ("file_size", "一条成片大约多大？", "500 MB 左右")]:
    reg(ws, r, 3, f"network.{key}", q)
    for a, b in ((1, 2), (3, 4), (5, 6)): ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
    c = ws.cell(r, 1, q); c.font = font(); c.alignment = WRAP
    for col in range(1, 7): ws.cell(r, col).border = BOX
    for col in (3, 4): ws.cell(r, col).fill = FILL_IN
    ws.cell(r, 3).alignment = WRAP; ws.cell(r, 3).font = font()
    for col in (5, 6): ws.cell(r, col).fill = FILL_EX
    e = ws.cell(r, 5, ex); e.font = font(10, color="888888", italic=True); e.alignment = WRAP
    ws.row_dimensions[r].height = 30
    r += 1
ws.freeze_panes = "A4"

# ---------- 5 文件清单 ----------
ws = wb.create_sheet("4 文件清单"); setup(ws)
widths = [14, 46, 30, 14, 12]
for i, w in enumerate(widths): ws.column_dimensions[chr(65 + i)].width = w
title(ws, "四、请提供的文件", "下面是我们最需要的几类文件，已列好建议项。请填文件名、选提交方式，发出后在最后一列选“已提交”。原始文件请走网盘或 U 盘，不要走微信。", 5)
r = 4
header(ws, r, ["属于哪个需求", "文件内容", "文件名", "提交方式", "是否已提交"]); r += 1
example_row(ws, r, ["需求 1", "员工完整操作一遍的录屏", "发布录屏.mp4", "网盘", "已提交"]); r += 1
items = [
    ("每个需求", "员工从头到尾真实操作一遍的录屏"),
    ("每个需求", "这个需求的输入样例（比如一条要处理的原始视频、一份要录入的表格）"),
    ("每个需求", "这个需求的产出样例（比如一条做好的成片）"),
    ("每个需求", "相关软件的导出 / 设置页面截图（比如剪辑软件的导出设置）"),
    ("全部", "第三、五张表里提到的“关于”页、系统信息截图和设备照片"),
    ("全部", "现成的 SOP、操作手册、账号表模板（如有，账号可用代号）"),
]
f0 = r
input_rows(ws, r, len(items) + 6, 5, h=34)
for rr in range(r, r + len(items) + 6):
    for ci, key in enumerate(["need", "content", "filename", "method", "submitted"], 1):
        reg(ws, rr, ci, f"files.{rr-r}.{key}")
for i, (a, b) in enumerate(items):
    ws.cell(r + i, 1, a).font = font(); ws.cell(r + i, 2, b).font = font()
    for col in (1, 2): ws.cell(r + i, col).fill = PatternFill(None)
dv3 = DataValidation(type="list", formula1='"网盘,U 盘,其他"', allow_blank=True)
dv4 = DataValidation(type="list", formula1='"已提交,还没有,没有这个文件"', allow_blank=True)
ws.add_data_validation(dv3); ws.add_data_validation(dv4)
end = f0 + len(items) + 5
dv3.add(f"D{f0}:D{end}"); dv4.add(f"E{f0}:E{end}")
ws.freeze_panes = "A5"


# ---------- 6 特殊设备 ----------
ws = wb.create_sheet("5 特殊设备"); setup(ws)
ws.column_dimensions["A"].width = 30; ws.column_dimensions["B"].width = 24
for col in "CDEF": ws.column_dimensions[col].width = 22
title(ws, "五、特殊设备（没有可跳过）", "如果工作要用到普通电脑以外的设备，比如工控一体机、触摸屏、收银机、平板或大屏、自助终端、PLC、扫码枪，每台设备填一列。看不懂的地方拍照就行：铭牌、系统“关于”页、接口、平时操作的界面。", 6)
header(ws, 3, ["项目", "示例", "设备 1", "设备 2", "设备 3", "设备 4"])
DEV = [
    ("name", "设备叫法 / 用途", "1 号产线触摸屏", None),
    ("type", "设备类型", "触摸屏", '"工控一体机,触摸屏,收银机,安卓平板或大屏,自助终端,PLC 或控制器,扫码枪或打印机,其他"'),
    ("model", "品牌和型号（看铭牌）", "某品牌 XX-100", None),
    ("os", "系统（看“关于”或“设置”页，不知道就拍照）", "看起来像安卓", None),
    ("can_install", "能不能自己装软件", "不清楚", '"能,不能,不清楚"'),
    ("vendor", "厂商或维保方，怎么联系", "设备厂商，有维保合同，售后电话", None),
    ("warranty", "改动会不会影响保修", "不清楚", '"会,不会,不清楚"'),
    ("network", "联网情况", "只连公司内网", '"能上网,只连公司内网,不联网,不清楚"'),
    ("owner", "公司里谁负责这台设备", "设备科老陈", None),
    ("vendor_onsite", "厂商工程师能不能到场配合", "能，要提前一周约", None),
    ("downtime", "什么时候可以停下来调试", "每周日停产", None),
    ("interfaces", "有哪些接口（网口、串口、USB 等，拍照即可）", "1 个网口、2 个 USB", None),
    ("data_flow", "数据怎么进出（导出文件、共享文件夹、上位机软件、只能看屏幕……）", "每天导出一个表格到 U 盘", None),
    ("photos", "照片文件名（铭牌、关于页、接口、操作界面）", "触摸屏-铭牌.jpg、触摸屏-接口.jpg", None),
]
r = 4
for key, label, ex, options in DEV:
    a = ws.cell(r, 1, label); a.font = font(); a.alignment = WRAP; a.border = BOX
    e = ws.cell(r, 2, ex); e.font = font(10, color="888888", italic=True); e.fill = FILL_EX; e.alignment = WRAP; e.border = BOX
    for col in range(3, 7):
        c = ws.cell(r, col); c.fill = FILL_IN; c.border = BOX; c.alignment = WRAP; c.font = font()
        reg(ws, r, col, f"devices.{col-3}.{key}", label)
    if options:
        dv = DataValidation(type="list", formula1=options, allow_blank=True)
        ws.add_data_validation(dv); dv.add(f"C{r}:F{r}")
    ws.row_dimensions[r].height = 40
    r += 1
ws.freeze_panes = "C4"

import json, sys
import os
out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "fde-pre-visit-intake", "template", "intake_form.xlsx")
wb.save(out)
FILES_PRESET = [{"need": a, "content": b} for a, b in items]
json.dump({"template": os.path.basename(out), "files_preset": FILES_PRESET, "fields": FIELDS},
          open(os.path.join(os.path.dirname(os.path.abspath(out)), "fields.json"), "w", encoding="utf-8"),
          ensure_ascii=False, indent=1)
print(out, len(FIELDS), "fields")
